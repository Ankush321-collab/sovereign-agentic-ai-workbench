import pytest
from pathlib import Path
import docx
from backend.tools.psu_note_generator import generate_psu_approval_note

def test_psu_note_generation():
    """TDD Seam: Verify generated PSU Green-Sheet note contains official Secretariat structure."""
    test_data = {
        "file_reference_no": "IOCL/RHQ/PL-MAINT/2026/FL-402",
        "subject": "REPLACEMENT & SHUTDOWN APPROVAL FOR FLANGE FL-402 DUE TO CRITICAL WALL THINNING",
        "equipment_tag": "FL-402",
        "nominal_thickness_mm": 6.4,
        "measured_thickness_mm": 3.8,
        "retirement_thickness_mm": 4.2,
        "corrosion_rate_mm_year": 0.45,
        "remaining_life_years": -0.89,
        "compliance_status": "NON-COMPLIANT (CRITICAL RISK)",
        "statutory_standard": "ASME B31.3 Section 304.1.2 & OISD-105",
        "operational_risk": "High risk of containment loss and volatile hydrocarbon leak at 142 PSI operating pressure.",
        "financial_estimate_inr": "₹ 14,50,000 (Fourteen Lakhs Fifty Thousand Only)",
        "recommendation": "Immediate emergency derating to 60 PSI and scheduled replacement during upcoming maintenance outage."
    }

    result = generate_psu_approval_note(test_data, output_filename="Test_Approval_Note.docx")

    assert result["success"] is True
    file_path = Path(result["file_path"])
    assert file_path.exists()
    assert file_path.stat().st_size > 5000

    # Inspect document contents across paragraphs and table cells
    doc = docx.Document(str(file_path))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    full_text = "\n".join(texts)

    # Verify official PSU Secretariat markers
    assert "INDIAN OIL CORPORATION LIMITED" in full_text
    assert "NOTE SHEET FOR APPROVAL" in full_text
    assert "IOCL/RHQ/PL-MAINT/2026/FL-402" in full_text
    assert "FL-402" in full_text
    assert "ASME B31.3" in full_text
    assert "Fourteen Lakhs Fifty Thousand" in full_text

    # Verify tables existence (metadata, inspection telemetry, and signature blocks)
    assert len(doc.tables) >= 3
