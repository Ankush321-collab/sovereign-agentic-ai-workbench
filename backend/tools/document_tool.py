from pathlib import Path
import logging
from backend.config import OUTPUTS_DIR
from backend.services.multimodal_service import MultimodalService

logger = logging.getLogger("document_tool")

def generate_docx(title: str, content: str, output_filename: str = "Approval_Note.docx") -> dict:
    """Generates an executive/industrial Word (.docx) document deliverable."""
    output_path = OUTPUTS_DIR / output_filename
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        
        doc = docx.Document()
        
        # Heading
        heading = doc.add_heading(title, level=1)
        
        # Subheading / Metadata
        meta_para = doc.add_paragraph()
        r = meta_para.add_run("Sovereign AI Workbench — Official Confidential Deliverable\n")
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor(100, 100, 100)
        
        # Body Content
        for paragraph_text in content.split("\n\n"):
            if paragraph_text.strip():
                doc.add_paragraph(paragraph_text.strip())
                
        doc.save(str(output_path))
        return {
            "success": True,
            "filename": output_path.name,
            "file_path": str(output_path),
            "message": f"Successfully generated DOCX document: {output_path.name}"
        }
    except Exception as e:
        logger.warning(f"python-docx generation failed ({e}), writing fallback text file.")
        txt_path = output_path.with_suffix(".txt")
        txt_path.write_text(f"TITLE: {title}\n\n{content}", encoding="utf-8")
        return {
            "success": True,
            "filename": txt_path.name,
            "file_path": str(txt_path),
            "message": f"Generated Text fallback document: {txt_path.name}"
        }

def generate_pptx(title: str, slide_titles: list[str], slide_contents: list[str], output_filename: str = "Report.pptx") -> dict:
    """Generates a PowerPoint (.pptx) presentation deck."""
    output_path = OUTPUTS_DIR / output_filename
    try:
        import pptx
        prs = pptx.Presentation()
        
        # Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Sovereign AI Workbench // Air-Gapped Executive Presentation"
        
        # Content Slides
        bullet_slide_layout = prs.slide_layouts[1]
        for t, c in zip(slide_titles, slide_contents):
            s = prs.slides.add_slide(bullet_slide_layout)
            s.shapes.title.text = t
            tf = s.placeholders[1].text_frame
            tf.text = c
            
        prs.save(str(output_path))
        return {
            "success": True,
            "filename": output_path.name,
            "file_path": str(output_path),
            "message": f"Successfully generated PPTX document: {output_path.name}"
        }
    except Exception as e:
        logger.warning(f"python-pptx generation failed ({e}), writing fallback markdown.")
        md_path = output_path.with_suffix(".md")
        md_content = f"# {title}\n\n"
        for t, c in zip(slide_titles, slide_contents):
            md_content += f"## {t}\n{c}\n\n"
        md_path.write_text(md_content, encoding="utf-8")
        return {
            "success": True,
            "filename": md_path.name,
            "file_path": str(md_path),
            "message": f"Generated Markdown presentation fallback: {md_path.name}"
        }

def edit_spreadsheet(rows: list[list], output_filename: str = "Calculation.xlsx") -> dict:
    """Generates or updates a styled Excel (.xlsx) spreadsheet deliverable."""
    try:
        from backend.tools.spreadsheet_generator import generate_asme_calculation_workbook
        return generate_asme_calculation_workbook(rows, output_filename=output_filename)
    except Exception as ex:
        logger.warning(f"openpyxl generation failed ({ex}), falling back to standard CSV.")
        output_path = OUTPUTS_DIR / output_filename
        csv_path = output_path.with_suffix(".csv")
        import csv
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)
        return {
            "success": True,
            "filename": csv_path.name,
            "file_path": str(csv_path),
            "message": f"Generated CSV fallback spreadsheet: {csv_path.name}"
        }

async def ocr_document(file_path: str) -> dict:
    """Invokes Pankaj's Multimodal Pipeline for document OCR and visual inspection."""
    return await MultimodalService.process_document(file_path)
