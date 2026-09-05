import os
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import OUTPUTS_DIR

def generate_asme_calculation_workbook(rows: List[List[Any]], title: str = "ASME B31.3 Engineering Calculation & Asset Checklist", output_filename: str = "Calculation.xlsx") -> Dict[str, Any]:
    """
    Generates a beautifully styled, formula-capable Excel workbook (.xlsx)
    for ASME B31.3 plant calculations and P&ID equipment checklists.
    """
    os.makedirs(OUTPUTS_DIR, exist_ok=True)
    file_path = OUTPUTS_DIR / output_filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Integrity Calculation"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Title Banner (Row 1-2)
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="104C8C", end_color="104C8C", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Subtitle / Classification Banner (Row 2)
    ws.merge_cells("A2:E2")
    sub_cell = ws["A2"]
    sub_cell.value = "INDIAN OIL CORPORATION LIMITED // SOVEREIGN AIR-GAPPED WORKBENCH // CONFIDENTIAL"
    sub_cell.font = Font(name="Arial", size=8.5, bold=True, color="D0E0F0")
    sub_cell.fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Empty spacer row
    ws.row_dimensions[3].height = 10

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    start_row = 4
    if rows:
        headers = rows[0]
        # Write Headers
        ws.row_dimensions[start_row].height = 24
        for c_idx, h_val in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=c_idx)
            cell.value = str(h_val)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Write Data Rows
        for r_idx, r_data in enumerate(rows[1:], start_row + 1):
            ws.row_dimensions[r_idx].height = 20
            is_alt = (r_idx % 2 == 0)
            bg_color = "F9FAFB" if is_alt else "FFFFFF"

            for c_idx, val in enumerate(r_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = Font(name="Arial", size=9.5)
                cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Conditional styling for status columns
                val_str = str(val).upper()
                if any(kw in val_str for kw in ["CRITICAL", "VIOLATION", "REPLACE", "TRIP"]):
                    cell.font = Font(name="Arial", size=9.5, bold=True, color="B91C1C")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif any(kw in val_str for kw in ["SAFE", "NORMAL", "OPERATIONAL", "ACTIVE"]):
                    cell.font = Font(name="Arial", size=9.5, bold=True, color="047857")
                    cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(str(file_path))

    return {
        "success": True,
        "filename": output_filename,
        "file_path": str(file_path),
        "file_size": file_path.stat().st_size
    }
