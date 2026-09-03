import logging
from pathlib import Path
from typing import Optional, Dict, Any

logger = logging.getLogger("multimodal.converter")

# Check Microsoft MarkItDown availability
try:
    from markitdown import MarkItDown
    HAS_MARKITDOWN = True
except ImportError:
    HAS_MARKITDOWN = False
    logger.info("Microsoft MarkItDown not installed in environment, using native modular fallback converters.")

try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    import pandas as pd
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False


class DocumentConverter:
    """Multi-format document to structured Markdown converter using MarkItDown and native fallbacks."""

    def __init__(self):
        self._markitdown = MarkItDown() if HAS_MARKITDOWN else None

    def convert_to_markdown(self, file_path: str) -> str:
        """Convert a document (PDF, Word, Excel, PowerPoint, Text, Image) to clean Markdown."""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Cannot convert non-existent file: {file_path}")

        suffix = path.suffix.lower()

        # 1. Try Microsoft MarkItDown first
        if self._markitdown is not None:
            try:
                result = self._markitdown.convert(str(path))
                if result and hasattr(result, "text_content") and result.text_content.strip():
                    return result.text_content
            except Exception as e:
                logger.warning(f"MarkItDown conversion failed for {path.name}, falling back to native engine: {e}")

        # 2. Modular native converters based on format
        if suffix == ".pdf":
            return self._convert_pdf(path)
        elif suffix in [".docx", ".doc"]:
            return self._convert_docx(path)
        elif suffix in [".xlsx", ".xls", ".csv"]:
            return self._convert_spreadsheet(path)
        elif suffix in [".txt", ".md", ".json", ".log"]:
            return self._convert_plain_text(path)
        else:
            # For images, return filename header (detailed OCR is handled by OCREngine)
            return f"![{path.name}]({path.name})"

    def _convert_pdf(self, path: Path) -> str:
        """Extract text and tables from PDF as structured Markdown."""
        pages_content = []

        if HAS_PDFPLUMBER:
            try:
                with pdfplumber.open(str(path)) as pdf:
                    for i, page in enumerate(pdf.pages):
                        page_num = i + 1
                        page_text = page.extract_text() or ""
                        tables = page.extract_tables()

                        page_md = [f"## Page {page_num}\n"]
                        if page_text:
                            page_md.append(page_text)

                        for table in tables:
                            if table and len(table) > 1:
                                page_md.append("\n" + self._format_table_as_markdown(table) + "\n")

                        pages_content.append("\n".join(page_md))
                if pages_content and any(p.strip() for p in pages_content):
                    return "\n\n---\n\n".join(pages_content)
            except Exception as e:
                logger.debug(f"pdfplumber extraction note: {e}")

        if HAS_PYMUPDF:
            try:
                doc = fitz.open(str(path))
                for i, page in enumerate(doc):
                    text = page.get_text()
                    if text.strip():
                        pages_content.append(f"## Page {i + 1}\n\n{text}")
                doc.close()
                if pages_content:
                    return "\n\n---\n\n".join(pages_content)
            except Exception as e:
                logger.debug(f"PyMuPDF extraction note: {e}")

        return f"# Document: {path.name}\n\n[Scanned document content - OCR required]"

    def _convert_docx(self, path: Path) -> str:
        if not HAS_DOCX:
            return f"# Document: {path.name}\n\n[python-docx required for Word parsing]"
        doc = docx.Document(str(path))
        lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                if p.style and "heading" in p.style.name.lower():
                    lines.append(f"### {p.text}")
                else:
                    lines.append(p.text)
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                table_data.append([cell.text.strip() for cell in row.cells])
            if table_data:
                lines.append("\n" + self._format_table_as_markdown(table_data) + "\n")
        return "\n\n".join(lines)

    def _convert_spreadsheet(self, path: Path) -> str:
        if not HAS_EXCEL:
            return f"# Spreadsheet: {path.name}\n\n[pandas/openpyxl required for Excel parsing]"
        if path.suffix.lower() == ".csv":
            df = pd.read_csv(str(path))
            return df.to_markdown(index=False)
        else:
            excel_file = pd.ExcelFile(str(path))
            sheets_md = []
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheets_md.append(f"### Sheet: {sheet_name}\n\n{df.to_markdown(index=False)}")
            return "\n\n---\n\n".join(sheets_md)

    def _convert_plain_text(self, path: Path) -> str:
        with open(str(path), "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def _format_table_as_markdown(table_rows: list) -> str:
        if not table_rows or len(table_rows) < 1:
            return ""
        header = table_rows[0]
        sanitized_header = [str(col).replace("\n", " ").strip() or " " for col in header]
        separator = ["---"] * len(sanitized_header)

        md_rows = [
            "| " + " | ".join(sanitized_header) + " |",
            "| " + " | ".join(separator) + " |"
        ]

        for row in table_rows[1:]:
            # match column count
            sanitized_row = [str(cell).replace("\n", " ").strip() if cell is not None else "" for cell in row]
            while len(sanitized_row) < len(sanitized_header):
                sanitized_row.append("")
            md_rows.append("| " + " | ".join(sanitized_row[:len(sanitized_header)]) + " |")

        return "\n".join(md_rows)
